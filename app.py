import os, re, json
from pathlib import Path
import streamlit as st
from pypdf import PdfReader
from docx import Document
from openai import OpenAI

st.set_page_config(page_title="AI Resume Reviewer", page_icon="🧠", layout="wide")
st.markdown("""<style>
.block-container{max-width:1150px;padding-top:2rem}
.hero{padding:2rem 2.3rem;border-radius:24px;background:linear-gradient(135deg,#07111d,#10364a,#08708a);color:white;margin-bottom:1.3rem}
.hero h1{font-size:3rem;margin:0}.hero p{color:#c9e9f0}
.metric{border:1px solid #d9e2ec;border-radius:16px;padding:1rem;text-align:center;background:#f8fbfd}
.metric b{display:block;font-size:2rem}.metric span{font-size:.78rem;color:#607080}
</style>""", unsafe_allow_html=True)

def extract(f):
    ext=Path(f.name).suffix.lower()
    if ext==".pdf":
        return "\n".join(p.extract_text() or "" for p in PdfReader(f).pages)
    if ext==".docx":
        d=Document(f); out=[p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows: out.append(" | ".join(c.text for c in r.cells))
        return "\n".join(out)
    return f.getvalue().decode("utf-8","ignore")

stop=set("""the and for with from that this your you are was were will have has had our their they them into over under about using use used can may must should job role work team experience skills responsible including such through across more than all who what how where when why a an of to in on at as by or is be it we i""".split())
def words(s): return [w for w in re.findall(r"[A-Za-z][A-Za-z0-9+#.-]{1,}",s.lower()) if w not in stop and len(w)>2]

def local_review(resume, job):
    from collections import Counter
    rf=set(words(resume)); jf=Counter(words(job))
    keys=[x for x,_ in jf.most_common(50)]
    matched=[x for x in keys if x in rf]; missing=[x for x in keys if x not in rf][:18]
    actions={"built","developed","designed","implemented","created","led","improved","automated","analyzed","optimized","delivered","launched","deployed","programmed","managed","engineered","tested","integrated"}
    action=sum(w in actions for w in words(resume)); nums=len(re.findall(r"\b\d+(?:\.\d+)?%?\b",resume))
    sections=sum(x in resume.lower() for x in ["education","experience","projects","skills","certifications"])
    keyword=round(100*len(matched)/max(1,min(len(keys),30)))
    impact=min(100,35+action*4+min(nums,8)*6)
    structure=min(100,sections*20+10)
    overall=round(.45*keyword+.30*impact+.25*structure)
    return overall,keyword,round(impact),round(structure),matched[:20],missing,nums,action

def ai_review(resume,job,model):
    key=os.getenv("OPENAI_API_KEY")
    if not key: return None,"OPENAI_API_KEY is not set."
    prompt=f"""Act as a technical recruiter reviewing a resume for a job. Do not invent experience.
Return ONLY JSON with: summary, strengths, gaps, keywords_to_consider, bullet_rewrites,
project_ideas, interview_topics, ats_notes. bullet_rewrites must contain original, rewrite, reason.
RESUME:
{resume[:30000]}
JOB:
{job[:30000]}"""
    try:
        r=OpenAI(api_key=key).responses.create(model=model,input=prompt)
        return json.loads(r.output_text),""
    except Exception as e: return None,str(e)

st.markdown("""<div class="hero"><h1>🧠 AI Resume Reviewer</h1>
<p>Upload a resume, paste a job description, and get a practical recruiter-style review of keywords,
impact, gaps, bullet points, and interview preparation.</p></div>""",unsafe_allow_html=True)

with st.sidebar:
    st.header("Settings")
    model=st.text_input("AI model","gpt-5.6-luna")
    st.caption("Local analysis works without an API key. AI review uses OPENAI_API_KEY.")
    st.caption("Never commit an API key to GitHub.")

c1,c2=st.columns(2)
with c1: rf=st.file_uploader("1. Upload resume",type=["pdf","docx","txt","md"])
with c2: job=st.text_area("2. Paste job description",height=230)

if rf and job.strip():
    resume=extract(rf)
    st.success(f"Loaded {rf.name} • {len(resume.split()):,} words")
    if st.button("🔎 Review Resume",type="primary",use_container_width=True):
        overall,keyword,impact,structure,matched,missing,nums,actions=local_review(resume,job)
        st.subheader("Resume health check")
        cols=st.columns(4)
        for col,n,label in zip(cols,[overall,keyword,impact,structure],["Overall heuristic","Keyword match","Impact signals","Structure"]):
            col.markdown(f'<div class="metric"><b>{n}</b><span>{label}</span></div>',unsafe_allow_html=True)
        st.caption("Heuristic signals only — not a real employer ATS score or an interview prediction.")
        a,b=st.columns(2)
        with a:
            st.markdown("### Keywords found")
            st.write(", ".join(matched) or "None detected")
        with b:
            st.markdown("### Keywords to inspect")
            st.write(", ".join(missing) or "No obvious gaps")
        if nums<3: st.warning("Consider adding truthful measurable evidence: users, $, %, counts, time saved, performance, or scope.")
        if actions<5: st.warning("Consider stronger action verbs at the start of bullets.")
        st.divider()
        st.subheader("🤖 AI recruiter review")
        data,err=ai_review(resume,job,model)
        if err:
            st.info(err+" Local analysis is still available above.")
        else:
            st.markdown("**Summary:** "+data.get("summary",""))
            x,y=st.columns(2)
            with x:
                st.markdown("**Strengths**")
                for v in data.get("strengths",[]): st.markdown("- "+v)
                st.markdown("**Gaps**")
                for v in data.get("gaps",[]): st.markdown("- "+v)
            with y:
                st.markdown("**Keywords to consider**")
                for v in data.get("keywords_to_consider",[]): st.markdown(f"- `{v}`")
                st.markdown("**Interview topics**")
                for v in data.get("interview_topics",[]): st.markdown("- "+v)
            st.markdown("### Bullet rewrites")
            for item in data.get("bullet_rewrites",[]):
                with st.expander(item.get("original","Original")):
                    st.markdown("**Rewrite:** "+item.get("rewrite",""))
                    st.caption(item.get("reason",""))
            st.markdown("### Project ideas")
            for v in data.get("project_ideas",[]): st.markdown("- "+v)
            st.markdown("### ATS notes")
            for v in data.get("ats_notes",[]): st.markdown("- "+v)
else:
    st.info("Upload a resume and paste a job description to start.")
    st.markdown("### V2 roadmap\n- PDF report export\n- Keyword highlighting\n- Resume version comparison\n- Job history\n- Portfolio/GitHub project recommendations\n- Unit tests and deployment")
